"""
Script to extract official dataset from Excel into TypeScript data modules.
ScamLens AI - Baseline 2026-10-27
"""

import openpyxl
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

EXCEL_PATH = r"d:\University Study Subjects\PROJECTS\AI-Innovation-DongNai-2026\ScamLens-AI\docs\ScamLens_AI_Official_Dataset_2026-10-27.xlsx"
OUTPUT_DIR = r"d:\University Study Subjects\PROJECTS\AI-Innovation-DongNai-2026\ScamLens-AI\src\lib\data"

os.makedirs(OUTPUT_DIR, exist_ok=True)
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

def sheet_to_dicts(sheet_name):
    sheet = wb[sheet_name]
    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    # filter out None headers
    headers = [h for h in headers if h is not None]
    rows = []
    for r in range(2, sheet.max_row + 1):
        # check if row is empty
        values = [sheet.cell(r, c + 1).value for c in range(len(headers))]
        if any(v is not None for v in values):
            row_dict = {}
            for h, v in zip(headers, values):
                row_dict[h] = v if v is not None else ""
            rows.append(row_dict)
    return rows

# 1. 06_DAU_MOI_LIEN_HE -> official-contacts.ts
contacts = sheet_to_dicts("06_DAU_MOI_LIEN_HE")
contacts_ts = f"""// Dữ liệu đầu mối liên hệ chính thống từ Bộ Công an & Tỉnh Đồng Nai
// Trích xuất từ ScamLens_AI_Official_Dataset_2026-10-27.xlsx (Sheet 06_DAU_MOI_LIEN_HE)

export interface OfficialContact {{
  agencyId: string;
  region: string;
  agencyName: string;
  department: string;
  scope: string;
  phone: string;
  email: string;
  website: string;
  address: string;
  priority: string;
  usageAdvice: string;
  verifiedStatus: string;
  sourceUrl: string;
}}

export const officialContacts: OfficialContact[] = {json.dumps([{
    'agencyId': c.get('agency_id', ''),
    'region': c.get('khu_vuc', ''),
    'agencyName': c.get('co_quan', ''),
    'department': c.get('don_vi', ''),
    'scope': c.get('phu_trach', ''),
    'phone': c.get('dien_thoai', ''),
    'email': c.get('email', ''),
    'website': c.get('website', ''),
    'address': c.get('dia_chi', ''),
    'priority': c.get('do_uu_tien', ''),
    'usageAdvice': c.get('cach_su_dung', ''),
    'verifiedStatus': c.get('trang_thai_xac_minh', ''),
    'sourceUrl': c.get('source_url', ''),
} for c in contacts], ensure_ascii=False, indent=2)};
"""
with open(os.path.join(OUTPUT_DIR, "official-contacts.ts"), "w", encoding="utf-8") as f:
    f.write(contacts_ts)

# 2. 07_QUY_TRINH_UNG_PHO -> playbooks.ts
playbooks = sheet_to_dicts("07_QUY_TRINH_UNG_PHO")
playbooks_ts = f"""// 14 Quy trình ứng phó khẩn cấp chính thức
// Trích xuất từ ScamLens_AI_Official_Dataset_2026-10-27.xlsx (Sheet 07_QUY_TRINH_UNG_PHO)

export interface Playbook {{
  playbookId: string;
  situation: string;
  immediateAction: string;
  nextStep: string;
  agencyIds: string[];
  inputTypes: string[];
  caution: string;
  sourceUrl: string;
}}

export const emergencyPlaybooks: Playbook[] = {json.dumps([{
    'playbookId': p.get('playbook_id', ''),
    'situation': p.get('tinh_huong', ''),
    'immediateAction': p.get('hanh_dong_ngay', ''),
    'nextStep': p.get('buoc_tiep', ''),
    'agencyIds': [a.strip() for a in str(p.get('agency_ids', '')).split(';') if a.strip()],
    'inputTypes': [i.strip() for i in str(p.get('input_types', '')).split(';') if i.strip()],
    'caution': p.get('chong_khang_dinh_qua_muc', ''),
    'sourceUrl': p.get('nguon_tieu_bieu', ''),
} for p in playbooks], ensure_ascii=False, indent=2)};
"""
with open(os.path.join(OUTPUT_DIR, "playbooks.ts"), "w", encoding="utf-8") as f:
    f.write(playbooks_ts)

# 3. 08_KENH_XAC_MINH -> verification-channels.ts
channels = sheet_to_dicts("08_KENH_XAC_MINH")
channels_ts = f"""// Kênh xác minh và tra cứu độc lập chính thống (VNNIC, Tổng đài 156, ...)
// Trích xuất từ ScamLens_AI_Official_Dataset_2026-10-27.xlsx (Sheet 08_KENH_XAC_MINH)

export interface VerificationChannel {{
  channelId: string;
  name: string;
  topic: string;
  channelAddress: string;
  canVerify: string;
  cannotConclude: string;
  sourceId: string;
  inputTypes: string[];
  sourceUrl: string;
}}

export const verificationChannels: VerificationChannel[] = {json.dumps([{
    'channelId': ch.get('channel_id', ''),
    'name': ch.get('ten', ''),
    'topic': ch.get('chu_de', ''),
    'channelAddress': ch.get('dia_chi_kenh', ''),
    'canVerify': ch.get('co_the_xac_minh', ''),
    'cannotConclude': ch.get('khong_the_ket_luan', ''),
    'sourceId': ch.get('source_id', ''),
    'inputTypes': [i.strip() for i in str(ch.get('input_types', '')).split(';') if i.strip()],
    'sourceUrl': ch.get('source_url', ''),
} for ch in channels], ensure_ascii=False, indent=2)};
"""
with open(os.path.join(OUTPUT_DIR, "verification-channels.ts"), "w", encoding="utf-8") as f:
    f.write(channels_ts)

# 4. 05_VAN_BAN_PHAP_LUAT -> legal-basis.ts
laws = sheet_to_dicts("05_VAN_BAN_PHAP_LUAT")
laws_ts = f"""// Căn cứ pháp lý an ninh mạng và bảo vệ dữ liệu cá nhân 2025 - 2026
// Trích xuất từ ScamLens_AI_Official_Dataset_2026-10-27.xlsx (Sheet 05_VAN_BAN_PHAP_LUAT)

export interface LegalBasis {{
  lawId: string;
  code: string;
  title: string;
  authority: string;
  promulgationDate: string;
  effectiveDate: string;
  status2026: string;
  notes: string;
  featureGroup: string;
  sourceUrl: string;
}}

export const legalBases: LegalBasis[] = {json.dumps([{
    'lawId': l.get('law_id', ''),
    'code': l.get('so_hieu', ''),
    'title': l.get('ten', ''),
    'authority': l.get('co_quan', ''),
    'promulgationDate': str(l.get('ban_hanh', '')),
    'effectiveDate': str(l.get('hieu_luc_tu', '')),
    'status2026': l.get('trang_thai_2026', ''),
    'notes': l.get('luu_y', ''),
    'featureGroup': l.get('nhom_tinh_nang', ''),
    'sourceUrl': l.get('source_url', ''),
} for l in laws], ensure_ascii=False, indent=2)};
"""
with open(os.path.join(OUTPUT_DIR, "legal-basis.ts"), "w", encoding="utf-8") as f:
    f.write(laws_ts)

# 5. 03_KICH_BAN_TIN_HIEU -> scenario-taxonomy.ts
scenarios = sheet_to_dicts("03_KICH_BAN_TIN_HIEU")
scenarios_ts = f"""// 38 Kịch bản thủ đoạn và phân loại tín hiệu cảnh báo
// Trích xuất từ ScamLens_AI_Official_Dataset_2026-10-27.xlsx (Sheet 03_KICH_BAN_TIN_HIEU)

export interface ScenarioTaxonomy {{
  scenarioId: string;
  group: string;
  title: string;
  sourceId: string;
  sourceAuthority: string;
  signalKeywords: string;
  findingCategories: string[];
  inputTypes: string[];
  riskPriority: string;
  explanation: string;
  sourceUrl: string;
}}

export const scenarioTaxonomies: ScenarioTaxonomy[] = {json.dumps([{
    'scenarioId': s.get('scenario_id', ''),
    'group': s.get('nhom', ''),
    'title': s.get('ten_kich_ban_bien_soan', ''),
    'sourceId': s.get('ma_nguon', ''),
    'sourceAuthority': s.get('co_quan_nguon', ''),
    'signalKeywords': s.get('mau_tin_hieu_tham_khao', ''),
    'findingCategories': [c.strip() for c in str(s.get('finding_categories', '')).split(';') if c.strip()],
    'inputTypes': [i.strip() for i in str(s.get('input_types', '')).split(';') if i.strip()],
    'riskPriority': s.get('risk_priority_for_triage', ''),
    'explanation': s.get('dien_giai', ''),
    'sourceUrl': s.get('source_url', ''),
} for s in scenarios], ensure_ascii=False, indent=2)};
"""
with open(os.path.join(OUTPUT_DIR, "scenario-taxonomy.ts"), "w", encoding="utf-8") as f:
    f.write(scenarios_ts)

print("All TypeScript data modules generated successfully in:", OUTPUT_DIR)
